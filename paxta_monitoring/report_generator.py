"""
Hisobotlarni shakllantirish moduli:
Tuman nomi har bir joyda aniq va ko'zga tashlanadigan qilib ko'rsatiladi.
"""
import os
import csv
import re
from html import escape as html_escape
import hashlib
import urllib.request
import urllib.error
from urllib.parse import quote, urljoin, urlparse
from copy import deepcopy
from collections import defaultdict
from datetime import datetime
from config import API_BASE_URL

RED = "\033[91m"
GREEN = "\033[92m"
YELLOW = "\033[93m"
BLUE = "\033[94m"
MAGENTA = "\033[95m"
CYAN = "\033[96m"
BOLD = "\033[1m"
RESET = "\033[0m"

_CYRILLIC_MONTHS = (
    "январь", "февраль", "март", "апрель", "май", "июнь",
    "июль", "август", "сентябрь", "октябрь", "ноябрь", "декабрь",
)
_MEASUREMENT_WORDS = (
    "Биринчи", "Иккинчи", "Учинчи", "Тўртинчи", "Бешинчи",
    "Олтинчи", "Еттинчи", "Саккизинчи", "Тўққизинчи", "Ўнинчи",
    "Ўн биринчи", "Ўн иккинчи", "Ўн учинчи", "Ўн тўртинчи",
    "Ўн бешинчи", "Ўн олтинчи", "Ўн еттинчи", "Ўн саккизинчи",
    "Ўн тўққизинчи", "Йигирманчи",
)

def measurement_label(index: int) -> str:
    """O'lchov tartibini so'z bilan, ortiqcha izohsiz beradi."""
    if 0 <= index < len(_MEASUREMENT_WORDS):
        return _MEASUREMENT_WORDS[index]
    return f"{index + 1}-"

def to_cyrillic(value: str) -> str:
    """Lotincha API nomlarini hisobotdagi o'zbek kirill yozuviga o'tkazadi."""
    text = str(value or "").strip()
    text = re.sub(r"(?i)mchj", "МЧЖ", text)
    text = re.sub(r"(?i)\bfx\b", "Ф/Х", text)
    text = text.replace("Gʻ", "G'").replace("gʻ", "g'").replace("ʼ", "'").replace("’", "'")
    replacements = (
        ("sh", "ш"), ("ch", "ч"), ("yo", "ё"), ("yu", "ю"), ("ya", "я"),
        ("o'", "ў"), ("g'", "ғ"),
    )
    text = text.lower()
    for source, target in replacements:
        text = text.replace(source, target)
    table = str.maketrans({
        "a": "а", "b": "б", "d": "д", "e": "е", "f": "ф", "g": "г",
        "h": "ҳ", "i": "и", "j": "ж", "k": "к", "l": "л", "m": "м",
        "n": "н", "o": "о", "p": "п", "q": "қ", "r": "р", "s": "с",
        "t": "т", "u": "у", "v": "в", "x": "х", "y": "й", "z": "з",
    })
    text = text.translate(table)
    return text[:1].upper() + text[1:] if text else text

def clean_maskan_name(value: str) -> str:
    """Maskan nomidan texnik suffixlar va yilni olib tashlaydi."""
    text = str(value or "").strip()
    text = re.split(r"(?i)\bptm\b", text, maxsplit=1)[0].strip(" ,-")
    match = re.search(r"(?i)\bmchj\b", text)
    if match:
        text = text[:match.end()]
    return text.strip(" ,-")

def _image_values(value, key_path=""):
    if isinstance(value, dict):
        for key, child in value.items():
            yield from _image_values(child, f"{key_path}_{key}".lower())
    elif isinstance(value, list):
        for index, child in enumerate(value):
            yield from _image_values(child, f"{key_path}_{index}".lower())
    elif isinstance(value, str) and (
        value.startswith(("http://", "https://", "/"))
        or value.lower().endswith((".jpg", ".jpeg", ".png", ".webp"))
    ):
        if any(word in key_path for word in ("image", "photo", "picture", "file", "media", "url")):
            yield key_path, value


def _image_direction(key_path: str) -> str:
    """Rasm kalitidan uning kirish yoki chiqish yukiga tegishliligini aniqlaydi."""
    output_words = ("out", "exit", "checkout", "weighout", "chiq", "tashq")
    return "out" if any(word in key_path for word in output_words) else "in"


def _image_index(key_path: str) -> int:
    """API kalitidagi kamera indeksini oladi; ko'rsatilmasa tabiiy tartib saqlanadi."""
    match = re.search(
        r"(?:image|images|photo|photos|picture|pictures|camera|cam|file|media)"
        r"[^0-9]*([01])(?:\D|$)",
        key_path,
    )
    return int(match.group(1)) if match else 99

class ReportGenerator:
    def __init__(self, analysis_result: dict):
        self.res = analysis_result
        self.r1 = analysis_result["rule1_results"]
        self.r2 = analysis_result["rule2_results"]
        self.r3 = analysis_result.get("rule3_results", [])
        self.r4 = analysis_result.get("rule4_results", [])
        self.r5 = analysis_result.get("rule5_results", [])

    @staticmethod
    def _weight_differences(first, second):
        fields = {
            "brutto": ("fullweight_f", "fullweight"),
            "tara": ("taraweight_f", "taraweight"),
            "netto": ("netto_f", "netto"),
        }
        differences = {}
        for name, (numeric_key, value_key) in fields.items():
            first_value = first.get(numeric_key)
            second_value = second.get(numeric_key)
            if first_value is None or second_value is None:
                first_value = first.get(value_key)
                second_value = second.get(value_key)
            if first_value is None or second_value is None:
                return None
            differences[name] = round(abs(float(first_value) - float(second_value)), 1)
        return differences

    @staticmethod
    def _farmer_page_url(load):
        raw = load.get("raw", {})
        region_id = raw.get("_region_id")
        district_id = raw.get("_district_id")
        farmer_tin = raw.get("_farmer_tax_number") or load.get("farmer_tin")
        if not region_id or not district_id or not farmer_tin or farmer_tin == "-":
            return None
        path_parts = (region_id, district_id, farmer_tin)
        encoded_path = "/".join(quote(str(part), safe="") for part in path_parts)
        return f"https://hosil-qabuli.uz/cotton/farmers/{encoded_path}"

    @classmethod
    def _case_controls(cls, case_id, loads):
        checkbox = (
            f"<input class='case-checkbox' type='checkbox' "
            f"data-case-id='{html_escape(case_id, quote=True)}' "
            f"aria-label='Holatni tanlash'>"
        )
        links = []
        seen_urls = set()
        for load in loads:
            url = cls._farmer_page_url(load)
            if not url or url in seen_urls:
                continue
            seen_urls.add(url)
            farmer = html_escape(str(load.get("farmer_name") or "Fermer xo'jaligi"))
            truck = html_escape(str(load.get("truck") or ""))
            label = f"{farmer} sahifasi"
            if truck:
                label += f" ({truck})"
            safe_url = html_escape(url, quote=True)
            links.append(
                f"<a class='source-link' href='{safe_url}' target='_blank' "
                f"rel='noopener noreferrer'>🔗 {label}ni ochish</a>"
            )
        if not links:
            links.append("<span class='source-link-unavailable'>Manba havolasi mavjud emas</span>")
        return checkbox, "<div class='source-links'>" + "".join(links) + "</div>"

    @staticmethod
    def _sortable_netto(load):
        value = load.get("netto_f")
        if value is None:
            value = load.get("netto")
        try:
            return float(str(value).replace(" ", "").replace(",", "."))
        except (TypeError, ValueError):
            return ""

    @classmethod
    def _sort_attributes(cls, item, intermediate_count, representative_load):
        maskan = html_escape(str(item.get("maskan") or ""), quote=True)
        farmer = html_escape(
            str(representative_load.get("farmer_name") or ""), quote=True
        )
        netto = cls._sortable_netto(representative_load)
        return (
            f"data-sort-count='{intermediate_count}' "
            f"data-sort-maskan='{maskan}' "
            f"data-sort-farmer='{farmer}' "
            f"data-sort-netto='{netto}'"
        )

    @staticmethod
    def _load_touches_date(load, report_date):
        return any(
            value is not None and value.date() == report_date
            for value in (load.get("in_dt"), load.get("out_dt"))
        )

    @classmethod
    def _result_touches_date(cls, result, report_date):
        loads = []
        for key in ("main_load", "first_load", "second_load"):
            load = result.get(key)
            if load:
                loads.append(load)
        for key in ("interleaved_loads", "all_intermediate_loads", "summary_measurements"):
            loads.extend(result.get(key, ()))
        return any(cls._load_touches_date(load, report_date) for load in loads)

    def print_console_summary(self):
        print(f"\n{BOLD}{CYAN}================================================================================{RESET}")
        print(f"{BOLD}{CYAN}           PAXTA MONITORINGI: ANOMALIYALAR VA AYLANTIRISHLAR TAHLILI            {RESET}")
        print(f"{BOLD}{CYAN}================================================================================{RESET}")
        print(f"Jami tekshirilgan o'lchovlar: {BOLD}{self.res['total_loads']}{RESET} | Maskanlar soni: {BOLD}{self.res['total_maskans']}{RESET}")
        print(f"{BOLD}{RED}🚨 3-QOIDA (AYNI BIR TEXNIKA + ORALIQDA KIRISH - AYLANTIRISH): {self.res['rule3_count']} TA HOLAT!{RESET}")
        print(f"1-Qoida bo'yicha umumiy oraliqdagi kirishlar: {BOLD}{self.res['rule1_count']} ta{RESET}")
        print(f"2-Qoida bo'yicha 15 daqiqalik takroriy kirishlar: {BOLD}{self.res['rule2_count']} ta{RESET}")
        print(f"4-Qoida bo'yicha og'irlik farqi va oraliq kirishlar: {BOLD}{self.res.get('rule4_count', 0)} ta{RESET}")
        print(f"5-Qoida bo'yicha mos fermer va og'irlikdagi oraliq holatlar: {BOLD}{self.res.get('rule5_count', 0)} ta{RESET}")
        print(f"{CYAN}--------------------------------------------------------------------------------{RESET}")

        # 3-QOIDA NATIJALARI
        print(f"\n{BOLD}{RED}🔴 [3-QOIDA] AYNI BIR TEXNIKA CHIQMASDAN TURIB QAYTA KIRGAN HOLATLAR ({len(self.r3)} ta):{RESET}")
        if not self.r3:
            print(f"  {GREEN}Bunday o'ta jiddiy qoidabuzarlik topilmadi.{RESET}")
        else:
            for idx, item in enumerate(self.r3, start=1):
                f1 = item["first_load"]
                f2 = item["second_load"]
                print(f"\n{BOLD}{RED}🚨 XAVF DARAJASI: O'TA YUQORI (AYLANTIRISH!){RESET}")
                print(f"   📍 TUMAN: {BOLD}{GREEN}{item['district'].upper()}{RESET}")
                print(f"   🏢 Maskan: {BOLD}{item['maskan']}{RESET}")
                print(f"   🚜 Texnika davlat raqami: {BOLD}{YELLOW}[{item['truck']}]{RESET}")
                print(f"   -------------------------------------------------------------------------")
                print(f"   1-Yuk (Asosiy):")
                print(f"       Tuman: {BOLD}{f1['district']}{RESET} | Fermer: {f1['farmer_name']} (INN: {f1['farmer_tin']})")
                print(f"       Kirish: {BOLD}{f1['in_str']}{RESET} -> Chiqish: {BOLD}{f1['out_str']}{RESET} | Netto: {f1['netto']} kg (Reg: № {f1['registration_number']})")
                print(f"   2-Yuk (Oraliqda kiritilgan):")
                print(f"       Tuman: {BOLD}{f2['district']}{RESET} | Fermer: {f2['farmer_name']} (INN: {f2['farmer_tin']})")
                print(f"       Kirish: {BOLD}{RED}{f2['in_str']}{RESET} -> Chiqish: {BOLD}{f2['out_str']}{RESET} | Netto: {f2['netto']} kg (Reg: № {f2['registration_number']})")
                print(f"   -------------------------------------------------------------------------")
                print(f"   {RED}XULOSA: [{item['truck']}] 1-yuk bo'yicha tarozi/maskandan chiqmasdan ({f1['out_str']}), xuddi shu texnika {f2['in_str']} da yana kirib boshqa fermerga hosil yozdirgan!{RESET}")

        # 1-QOIDA NATIJALARI
        print(f"\n{BOLD}{MAGENTA}[1-QOIDA] Traktor kirish va chiqish oralig'ida boshqalar kirishi ({len(self.r1)} ta):{RESET}")
        for idx, item in enumerate(self.r1[:6], start=1):
            m = item["main_load"]
            print(f"   {idx}. {BOLD}{GREEN}[{item['district']}]{RESET} {item['maskan']} | Texnika: [{m['truck']}] Kirish: {m['in_str']} -> Chiqish: {m['out_str']} | Oraliqda {item['count']} ta kirish")

        if len(self.r1) > 6:
            print(f"      ... va yana {len(self.r1) - 6} ta holat (barchasi CSV / HTML da)")

        # 2-QOIDA NATIJALARI
        print(f"\n{BOLD}{YELLOW}[2-QOIDA] 15 daqiqa ichida takroriy kirish o'tkazilgan holatlar ({len(self.r2)} ta):{RESET}")
        for idx, item in enumerate(self.r2[:6], start=1):
            f1 = item["first_load"]
            f2 = item["second_load"]
            diff = item["diff_minutes"]
            print(f"   {idx}. {BOLD}{GREEN}[{item['district']}]{RESET} {item['maskan']} | [{f1['truck']}] va [{f2['truck']}] | Farq: {diff} daqiqa")

        if len(self.r2) > 6:
            print(f"      ... va yana {len(self.r2) - 6} ta holat (barchasi CSV / HTML da)")

        print(f"\n{BOLD}{RED}[4-QOIDA] Brutto farqi 400 kg, tara 100 kg, netto 300 kg gacha va oraliqdagi boshqa yuklar ({len(self.r4)} ta):{RESET}")
        for idx, item in enumerate(self.r4[:6], start=1):
            first = item["first_load"]
            second = item["second_load"]
            print(
                f"   {idx}. {BOLD}{GREEN}[{item['district']}]{RESET} {item['maskan']} | "
                f"1-yuk: {first['truck']} ({first['farmer_name']}, {first['district']}) | "
                f"2-yuk: {second['truck']} ({second['farmer_name']}, {second['district']}) | "
                f"Brutto farqi: {item['weight_differences']['brutto']} kg | "
                f"Tara farqi: {item['weight_differences']['tara']} kg | "
                f"Netto farqi: {item['weight_differences']['netto']} kg"
            )
        if len(self.r4) > 6:
            print(f"      ... va yana {len(self.r4) - 6} ta holat (barchasi CSV / HTML da)")

        print(f"\n{BOLD}{MAGENTA}[5-QOIDA] Bir xil fermer va og'irlik shartiga mos oraliqdagi kirishlar ({len(self.r5)} ta):{RESET}")
        for idx, item in enumerate(self.r5[:6], start=1):
            main = item["main_load"]
            intermediate_trucks = ", ".join(
                f"{load['truck']} ({load['in_str']})"
                for load in item["interleaved_loads"]
            )
            print(
                f"   {idx}. {BOLD}{GREEN}[{item['district']}]{RESET} {item['maskan']} | "
                f"Asosiy texnika: [{main['truck']}] {main['in_str']} -> {main['out_str']} | "
                f"Oraliqda {item['count']} ta: {intermediate_trucks}"
            )
        if len(self.r5) > 6:
            print(f"      ... va yana {len(self.r5) - 6} ta holat (barchasi CSV / HTML da)")

    def export_csv(self, output_dir: str = "."):
        os.makedirs(output_dir, exist_ok=True)
        
        # 3-Qoida CSV (AYLANTIRISH)
        r3_path = os.path.join(output_dir, "hisobot_3_aylantirish_ayni_texnika.csv")
        with open(r3_path, "w", newline="", encoding="utf-8-sig") as fp:
            writer = csv.writer(fp)
            writer.writerow([
                "T/r", "Tuman", "Maskan", "Texnika Raqami",
                "1-Kirish", "1-Chiqish", "1-Netto (kg)", "1-Tuman", "1-Fermer", "1-INN", "1-Reg. №",
                "2-Kirish", "2-Chiqish", "2-Netto (kg)", "2-Tuman", "2-Fermer", "2-INN", "2-Reg. №",
                "Ustma-ust tushgan vaqt (daqiqa)"
            ])
            for idx, item in enumerate(self.r3, start=1):
                f1 = item["first_load"]
                f2 = item["second_load"]
                writer.writerow([
                    idx, item["district"], item["maskan"], item["truck"],
                    f1["in_str"], f1["out_str"], f1["netto"], f1["district"], f1["farmer_name"], f1["farmer_tin"], f1["registration_number"],
                    f2["in_str"], f2["out_str"], f2["netto"], f2["district"], f2["farmer_name"], f2["farmer_tin"], f2["registration_number"],
                    item["overlap_minutes"]
                ])

        # 1-Qoida CSV
        r1_path = os.path.join(output_dir, "hisobot_1_oraliq_kirishlar.csv")
        with open(r1_path, "w", newline="", encoding="utf-8-sig") as fp:
            writer = csv.writer(fp)
            writer.writerow([
                "T/r", "Tuman", "Maskan", 
                "Asosiy Texnika", "Asosiy Kirish", "Asosiy Chiqish", "Asosiy Netto (kg)", "Asosiy Fermer", "Asosiy Reg. №",
                "Oraliqdagi Texnika", "Oraliqdagi Kirish", "Oraliqdagi Chiqish", "Oraliqdagi Netto (kg)", "Oraliqdagi Tuman", "Oraliqdagi Fermer", "Oraliqdagi Reg. №"
            ])
            row_idx = 1
            for item in self.r1:
                m = item["main_load"]
                for sub in item["interleaved_loads"]:
                    writer.writerow([
                        row_idx, item["district"], item["maskan"],
                        m["truck"], m["in_str"], m["out_str"], m["netto"], m["farmer_name"], m["registration_number"],
                        sub["truck"], sub["in_str"], sub["out_str"], sub["netto"], sub["district"], sub["farmer_name"], sub["registration_number"]
                    ])
                    row_idx += 1

        # 2-Qoida CSV
        r2_path = os.path.join(output_dir, "hisobot_2_15min_kirishlar.csv")
        with open(r2_path, "w", newline="", encoding="utf-8-sig") as fp:
            writer = csv.writer(fp)
            writer.writerow([
                "T/r", "Tuman", "Maskan", "Bir xil texnikami?", "Vaqt farqi (daqiqa)",
                "1-Texnika", "1-Kirish", "1-Netto (kg)", "1-Tuman", "1-Fermer", "1-Reg. №",
                "2-Texnika", "2-Kirish", "2-Netto (kg)", "2-Tuman", "2-Fermer", "2-Reg. №"
            ])
            for idx, item in enumerate(self.r2, start=1):
                f1 = item["first_load"]
                f2 = item["second_load"]
                writer.writerow([
                    idx, item["district"], item["maskan"],
                    "HA (AYNI TEXNIKA)" if item["is_same_truck"] else "Yo'q",
                    item["diff_minutes"],
                    f1["truck"], f1["in_str"], f1["netto"], f1["district"], f1["farmer_name"], f1["registration_number"],
                    f2["truck"], f2["in_str"], f2["netto"], f2["district"], f2["farmer_name"], f2["registration_number"]
                ])

        r4_path = os.path.join(output_dir, "hisobot_4_ogirlik_oraliq_kirishlar.csv")
        with open(r4_path, "w", newline="", encoding="utf-8-sig") as fp:
            writer = csv.writer(fp)
            writer.writerow([
                "T/r", "Tuman", "Maskan", "Brutto farqi (kg)", "Tara farqi (kg)", "Netto farqi (kg)",
                "1-Texnika", "1-Fermer", "1-Fermer tumani", "1-Kirish", "1-Chiqish", "1-Brutto", "1-Tara", "1-Netto",
                "2-Texnika", "2-Fermer", "2-Fermer tumani", "2-Kirish", "2-Chiqish", "2-Reg. №",
            ])
            for idx, item in enumerate(self.r4, start=1):
                first = item["first_load"]
                second = item["second_load"]
                writer.writerow([
                    idx, item["district"], item["maskan"],
                    item["weight_differences"]["brutto"],
                    item["weight_differences"]["tara"],
                    item["weight_differences"]["netto"],
                    first["truck"], first["farmer_name"], first["district"], first["in_str"], first["out_str"],
                    first["fullweight"], first["taraweight"], first["netto"],
                    second["truck"], second["farmer_name"], second["district"], second["in_str"], second["out_str"],
                    second["registration_number"],
                ])

        print(f"\n{GREEN}CSV hisobotlar saqlandi:{RESET}")
        print(f"  ⭐ {r3_path} (AYLANTIRISHLAR)")
        print(f"  1. {r1_path}")
        print(f"  2. {r2_path}")
        print(f"  4. {r4_path}")
        r5_path = os.path.join(output_dir, "hisobot_5_bitta_fermer_kirishlar.csv")
        with open(r5_path, "w", newline="", encoding="utf-8-sig") as fp:
            writer = csv.writer(fp)
            writer.writerow([
                "T/r", "Tuman", "Maskan",
                "Asosiy texnika", "Asosiy kirish", "Asosiy chiqish", "Asosiy netto (kg)",
                "Asosiy fermer", "Asosiy fermer tumani", "Oraliqdagi texnika",
                "Oraliqdagi kirish", "Oraliqdagi chiqish", "Oraliqdagi netto (kg)",
                "Oraliqdagi fermer", "Oraliqdagi fermer tumani", "Brutto farqi (kg)",
                "Tara farqi (kg)", "Netto farqi (kg)",
            ])
            row_idx = 1
            for item in self.r5:
                first = item["main_load"]
                for second in item["interleaved_loads"]:
                    differences = self._weight_differences(first, second) or {}
                    writer.writerow([
                        row_idx, item["district"], item["maskan"],
                        first["truck"], first["in_str"], first["out_str"], first["netto"],
                        first["farmer_name"], first["district"], second["truck"],
                        second["in_str"], second["out_str"], second["netto"],
                        second["farmer_name"], second["district"],
                        round(differences.get("brutto", 0), 1),
                        round(differences.get("tara", 0), 1),
                        round(differences.get("netto", 0), 1),
                    ])
                    row_idx += 1
        print(f"  5. {r5_path}")

    def export_pptx_reports(
        self,
        output_dir: str = ".",
        template_path: str = "",
        report_date=None,
        image_headers=None,
        one_report_per_case: bool = False,
    ):
        """Sana bo'yicha PPTX yaratadi; kerak bo'lsa har bir holatni alohida fayl qiladi."""
        try:
            from pptx import Presentation
            from pptx.dml.color import RGBColor
            from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
            from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
            from pptx.oxml.xmlchemy import OxmlElement
            from pptx.util import Inches, Pt
        except ImportError as exc:
            raise RuntimeError("PPTX eksporti uchun python-pptx o'rnatilishi kerak") from exc

        report_date = report_date or datetime.now().date()
        grouped = defaultdict(list)
        case_number = 0
        for item in self.r3:
            first_dt = item["first_load"].get("in_dt")
            if not first_dt or not self._result_touches_date(item, report_date):
                continue
            key = (item.get("district") or "Noma'lum tuman",
                   item["first_load"].get("farmer_name") or "Noma'lum fermer")
            if one_report_per_case:
                key = (*key, case_number)
                case_number += 1
            grouped[key].append(item)

        os.makedirs(output_dir, exist_ok=True)
        image_cache = os.path.join(output_dir, ".image_cache")
        os.makedirs(image_cache, exist_ok=True)
        def load_images(load):
            raw = load.get("raw", {})
            values = list(_image_values(raw))
            grouped = {"in": [], "out": []}
            for key, value in values:
                direction = _image_direction(key)
                if not any(value == existing[1] for existing in grouped[direction]):
                    grouped[direction].append((_image_index(key), value))
            result = {}
            for direction, image_values in grouped.items():
                image_values.sort(key=lambda item: item[0])
                urls = [url for _, url in image_values]
                if not urls:
                    continue

                if direction == "in":
                    # Kirishda o'rtadagi rasm old, oxirgi rasm orqa kamera hisoblanadi.
                    camera_urls = {
                        0: urls[len(urls) // 2],
                        1: urls[-1],
                    }
                else:
                    # Chiqishda birinchi rasm orqa, oxirgi rasm old kamera hisoblanadi.
                    camera_urls = {
                        0: urls[-1],
                        1: urls[0],
                    }

                for camera_index, url in camera_urls.items():
                    if url.startswith("/") and not os.path.exists(url):
                        url = urljoin(f"{API_BASE_URL.rstrip('/')}/", url.lstrip("/"))
                    if url.startswith(("http://", "https://")):
                        suffix = os.path.splitext(urlparse(url).path)[1] or ".jpg"
                        path = os.path.join(
                            image_cache,
                            hashlib.sha256(url.encode()).hexdigest()[:24]
                            + suffix,
                        )
                        if not os.path.exists(path):
                            try:
                                request = urllib.request.Request(url, headers=image_headers or {})
                                with urllib.request.urlopen(request, timeout=30) as response:
                                    with open(path, "wb") as output:
                                        output.write(response.read())
                            except (OSError, urllib.error.URLError, ValueError):
                                continue
                    else:
                        path = url
                    if os.path.exists(path):
                        result[(direction, camera_index)] = path
            return result

        def add_photo(slide, load, direction, camera_index, slot):
            if not load:
                return
            image = load_images(load).get((direction, camera_index))
            if not image:
                return
            picture_placeholders = sorted(
                (
                    shape for shape in slide.shapes
                    if shape.is_placeholder
                    and shape.placeholder_format.type == PP_PLACEHOLDER.PICTURE
                ),
                key=lambda shape: shape.top,
            )
            pictures = sorted(
                (
                    shape for shape in slide.shapes
                    if shape.element.tag.endswith("}pic")
                ),
                key=lambda shape: shape.top,
            )
            if slot < len(picture_placeholders):
                picture_placeholders[slot].insert_picture(image)
            elif slot < len(pictures):
                picture = pictures[slot]
                left, top, width, height = (
                    picture.left, picture.top, picture.width, picture.height
                )
                slide.shapes._spTree.remove(picture.element)
                slide.shapes.add_picture(image, left, top, width=width, height=height)

        created = []
        if template_path and os.path.exists(template_path):
            def load_sort_key(load):
                return (
                    load.get("in_dt") or datetime.max,
                    load.get("out_dt") or datetime.max,
                    str(load.get("id") or load.get("registration_number") or ""),
                )

            for report_number, (group_key, items) in enumerate(
                sorted(
                    grouped.items(),
                    key=lambda group: (
                        group[0][0], group[0][1],
                        group[0][2] if len(group[0]) > 2 else 0,
                    ),
                ),
                start=1,
            ):
                district, farmer = group_key[:2]
                items.sort(key=lambda x: (
                    x["first_load"]["in_dt"],
                    x["second_load"].get("in_dt") or x["first_load"]["in_dt"],
                ))
                item = items[0]
                first = item["first_load"]
                measurements = {}
                summary_measurements = {}
                primary_loads = {}
                intermediate_loads = {}
                measurement_numbers = {}

                def measurement_key(load):
                    return load.get("id") or (
                        load.get("in_dt"), load.get("out_dt"), load.get("registration_number")
                    )

                for entry in items:
                    for load in (entry["first_load"], entry["second_load"]):
                        measurements[measurement_key(load)] = load
                    for load in entry.get(
                        "summary_measurements",
                        (entry["first_load"], entry["second_load"]),
                    ):
                        summary_measurements[measurement_key(load)] = load
                    primary_loads[measurement_key(entry["first_load"])] = entry["first_load"]
                    for load in entry.get("all_intermediate_loads", (entry["second_load"],)):
                        intermediate_loads[measurement_key(load)] = load
                    for load, measurement_number in entry.get("measurement_numbers", ()):
                        measurement_numbers[measurement_key(load)] = measurement_number
                measurements = sorted(
                    measurements.values(),
                    key=load_sort_key,
                )
                summary_measurements = sorted(
                    summary_measurements.values(),
                    key=load_sort_key,
                )
                measurement_pairs = [
                    measurements[index:index + 2]
                    for index in range(0, len(measurements), 2)
                ]
                if not measurement_pairs:
                    continue
                final_out = max(
                    (
                        load.get("out_dt")
                        for load in summary_measurements
                        if load.get("out_dt")
                    ),
                    default=summary_measurements[-1].get("out_dt"),
                )
                first_dt = first["in_dt"]
                first_time = first_dt.strftime("%H:%M") if first_dt else "-"
                final_time = final_out.strftime("%H:%M") if final_out else "-"
                date_label = f"{first_dt.day}-{_CYRILLIC_MONTHS[first_dt.month - 1]}" if first_dt else "-"
                district_cyr = to_cyrillic(district)
                district_base_cyr = re.sub(r"(?i)\s+тумани$", "", district_cyr).strip()
                maskan_cyr = clean_maskan_name(first["maskan"]).upper()
                farmer_base = re.sub(r"(?i)\s+(fx|ф/х)$", "", str(farmer)).strip()
                farmer_cyr = farmer_base.upper()
                measurement_count = len(summary_measurements)
                primary_count = len(primary_loads)
                intermediate_count = len(intermediate_loads)
                total_netto_tons = sum(
                    load.get("netto_f") or 0
                    for load in summary_measurements
                ) / 1000
                total_netto_label = f"{total_netto_tons:.1f}".replace(".", ",")
                prs = Presentation(template_path)

                def replace_paragraph(paragraph, segments):
                    """Abzas formatini va har bir matn bo'lagining uslubini saqlab almashtiradi."""
                    original_runs = list(paragraph.runs)
                    if not original_runs:
                        return
                    style_rprs = [
                        deepcopy(run._r.get_or_add_rPr())
                        for run in original_runs
                    ]
                    style_fonts = [
                        (
                            run.font.name,
                            run.font.size,
                            run.font.bold,
                            run.font.italic,
                            run.font.underline,
                        )
                        for run in original_runs
                    ]
                    for child in list(paragraph._p):
                        if child.tag.endswith(("}r", "}br", "}fld")):
                            paragraph._p.remove(child)
                    for index, (text, style_index) in enumerate(segments):
                        pieces = str(text).split("\v")
                        for piece_index, piece in enumerate(pieces):
                            run = paragraph.add_run()
                            target_rpr = run._r.get_or_add_rPr()
                            for child in list(target_rpr):
                                target_rpr.remove(child)
                            for child in style_rprs[min(style_index, len(style_rprs) - 1)]:
                                target_rpr.append(deepcopy(child))
                            font_name, font_size, font_bold, font_italic, font_underline = (
                                style_fonts[min(style_index, len(style_fonts) - 1)]
                            )
                            run.font.name = font_name
                            run.font.size = font_size
                            run.font.bold = font_bold
                            run.font.italic = font_italic
                            run.font.underline = font_underline
                            run.text = piece
                            if piece_index < len(pieces) - 1:
                                run._r.addnext(OxmlElement("a:br"))

                def replace_paragraph_matches(paragraph, replacements):
                    """Matn shablonini almashtirib, o'zgarmagan qismlarning run formatlarini saqlaydi."""
                    original_text = paragraph.text
                    original_runs = list(paragraph.runs)
                    if not original_runs:
                        paragraph.text = original_text
                        return

                    char_styles = [0] * len(original_text)
                    cursor = 0
                    for style_index, run in enumerate(original_runs):
                        if not run.text:
                            continue
                        start = original_text.find(run.text, cursor)
                        if start < 0:
                            raise ValueError(
                                f"Shablon matnidagi run topilmadi: {run.text!r}"
                            )
                        for char_index in range(cursor, start):
                            char_styles[char_index] = style_index
                        for char_index in range(start, start + len(run.text)):
                            char_styles[char_index] = style_index
                        cursor = start + len(run.text)
                    for char_index in range(cursor, len(original_text)):
                        char_styles[char_index] = (
                            char_styles[char_index - 1] if char_index else 0
                        )

                    replacements = sorted(replacements, key=lambda item: item[0])
                    segments = []
                    cursor = 0

                    def append_original(start, end):
                        for char_index in range(start, end):
                            style_index = char_styles[char_index]
                            if segments and segments[-1][1] == style_index:
                                segments[-1] = (
                                    segments[-1][0] + original_text[char_index],
                                    style_index,
                                )
                            else:
                                segments.append((original_text[char_index], style_index))

                    for start, end, value in replacements:
                        if start < cursor:
                            raise ValueError("Shablon matni o'rinbosarlari ustma-ust tushdi")
                        append_original(cursor, start)
                        style_index = char_styles[start] if start < len(char_styles) else 0
                        segments.append((value, style_index))
                        cursor = end
                    append_original(cursor, len(original_text))
                    replace_paragraph(paragraph, segments)

                def replace_text(shape, value):
                    """Oddiy matnlar uchun mavjud birinchi run formatini saqlaydi."""
                    if not shape.has_text_frame:
                        return
                    paragraph = shape.text_frame.paragraphs[0]
                    if paragraph.runs:
                        replace_paragraph(paragraph, [(str(value), 0)])
                    else:
                        paragraph.text = str(value)

                def clone_slide(source_slide):
                    """Template slaydini rasmlari va foto joylari bilan nusxalaydi."""
                    target_slide = prs.slides.add_slide(source_slide.slide_layout)
                    for shape in list(target_slide.shapes):
                        target_slide.shapes._spTree.remove(shape.element)
                    rel_map = {}
                    for rel in source_slide.part.rels.values():
                        if rel.reltype.endswith("/slideLayout"):
                            continue
                        new_rel = target_slide.part.relate_to(
                            rel.target_ref if rel.is_external else rel.target_part,
                            rel.reltype,
                            is_external=rel.is_external,
                        )
                        rel_map[rel.rId] = new_rel
                    for shape in source_slide.shapes:
                        copied = deepcopy(shape.element)
                        for element in copied.iter():
                            for attribute, value in list(element.attrib.items()):
                                if value in rel_map:
                                    element.set(attribute, rel_map[value])
                        target_slide.shapes._spTree.insert_element_before(copied, "p:extLst")
                    return target_slide

                if prs.slides and prs.slides[0].shapes:
                    title_shape = prs.slides[0].shapes[0]
                    title_paragraphs = title_shape.text_frame.paragraphs
                    replace_paragraph(title_paragraphs[0], [(
                        f"Фарғона вилояти {district_cyr}да {first_dt.year} йил пахта йиғим-терими "
                        f"электрон тарози ўлчови тўғрисида таҳлилий\vМАЪЛУМОТ", 0
                    )])
                    summary_paragraph = title_paragraphs[1]
                    summary_text = summary_paragraph.text
                    replacements = []

                    def add_replacement(match, value, group=0):
                        start, end = match.span(group)
                        replacements.append((start, end, value))

                    district_match = re.search(
                        r"(?<=        ).+?(?= туманидаги)", summary_text
                    )
                    maskan_match = re.search(
                        r"“([^”]+)” пахта қабул қилиш масканларида", summary_text
                    )
                    farmer_match = re.search(
                        r"масканларида “([^”]+)” ФХ томонидан", summary_text
                    )
                    date_match = re.search(r"\b\d{1,2}-\S+", summary_text)
                    time_matches = list(re.finditer(r"\b\d{2}:\d{2}\b", summary_text))
                    equipment_match = re.search(
                        r"1 та трактор ва 2 та тиркамада", summary_text
                    )
                    trailer_match = re.search(r"\(тележка\)", summary_text)
                    count_match = re.search(
                        r"(?<=\v)\d+(?= маротаба)", summary_text
                    )
                    weight_match = re.search(r"(?<=\()\d+,\d(?= тонна\))", summary_text)

                    required_matches = (
                        district_match, maskan_match, farmer_match, date_match,
                        equipment_match, trailer_match, count_match, weight_match,
                    )
                    if any(match is None for match in required_matches) or len(time_matches) < 2:
                        raise ValueError(
                            "5-qoida shablonidagi matn kutilgan tuzilishga mos kelmadi"
                        )

                    add_replacement(district_match, district_base_cyr)
                    add_replacement(maskan_match, maskan_cyr, 1)
                    add_replacement(farmer_match, farmer_cyr, 1)
                    add_replacement(date_match, date_label)
                    add_replacement(time_matches[0], first_time)
                    add_replacement(time_matches[-1], final_time)
                    add_replacement(
                        equipment_match,
                        f"{primary_count} та трактор ва "
                        f"{intermediate_count} та тиркамада",
                    )
                    add_replacement(trailer_match, "(тележка)")
                    add_replacement(count_match, str(measurement_count))
                    add_replacement(weight_match, total_netto_label)
                    replace_paragraph_matches(summary_paragraph, replacements)

                template_photo_slides = list(prs.slides)[1:5]
                photo_slide_batches = []
                for source_slide in template_photo_slides:
                    slides = [source_slide]
                    slides.extend(
                        clone_slide(source_slide)
                        for _ in range(len(measurement_pairs) - 1)
                    )
                    photo_slide_batches.append(slides)

                cotton_slides = []
                trailer_slides = []
                for source_index, is_cotton, slide_collection in (
                    (0, True, cotton_slides),
                    (1, True, cotton_slides),
                    (2, False, trailer_slides),
                    (3, False, trailer_slides),
                ):
                    for item_index, pair in enumerate(measurement_pairs):
                        pair_loads = (
                            pair[0],
                            pair[1] if len(pair) > 1 else None,
                        )
                        pair_labels = tuple(
                            measurement_label(
                                measurement_numbers.get(
                                    measurement_key(load),
                                    item_index * 2 + slot,
                                )
                            ) if load else ""
                            for slot, load in enumerate(pair)
                        )
                        if len(pair_labels) == 1:
                            pair_labels = (*pair_labels, "")
                        layout_name = template_photo_slides[source_index].slide_layout.name.lower()
                        is_trailer = "tirkama" in layout_name
                        if not is_trailer:
                            pair_times = tuple(
                                load.get("in_dt") if load else None
                                for load in pair_loads
                            )
                        else:
                            pair_times = tuple(
                                load.get("out_dt") if load else None
                                for load in pair_loads
                            )
                        slide = photo_slide_batches[source_index][item_index]
                        slide_collection.append(slide)
                        side = "орқа томондан" if "orqa" in layout_name else "олд томондан"
                        direction = "out" if is_trailer else "in"
                        camera_index = 1 if "orqa" in layout_name else 0
                        add_photo(slide, pair_loads[0], direction, camera_index, 0)
                        add_photo(slide, pair_loads[1], direction, camera_index, 1)
                        for shape in slide.shapes:
                            text = getattr(shape, "text", "")
                            slot = 0 if shape.top < prs.slide_height / 2 else 1
                            value_dt = pair_times[slot]
                            placeholder_idx = (
                                shape.placeholder_format.idx if shape.is_placeholder else None
                            )
                            if placeholder_idx in (14, 15):
                                replace_text(shape, pair_labels[slot] if value_dt else "")
                            elif placeholder_idx in (16, 17):
                                replace_text(
                                    shape,
                                    f"{value_dt.day:02d}.{value_dt.month:02d}.{value_dt.year} й Соат "
                                    f"{value_dt.strftime('%H:%M')}" if value_dt else "",
                                )
                            elif ("12.09.2026" in text or "Соат" in text or "Soat" in text) and any(
                                char.isdigit() for char in text
                            ):
                                replace_text(
                                    shape,
                                    f"{value_dt.day:02d}.{value_dt.month:02d}.{value_dt.year} й Соат "
                                    f"{value_dt.strftime('%H:%M')}" if value_dt else "",
                                )
                            elif text in ("Биринчи ўлчов", "Иккинчи ўлчов"):
                                replace_text(shape, pair_labels[slot])
                            elif text in ("олд томондан", "орқа томондан"):
                                replace_text(shape, side if value_dt else "")

                # Har bir o'lchov juftligi o'zining 4 slaydlik template blokini saqlaydi:
                # paxta oldi, paxta orqasi, tirkama oldi, tirkama orqasi.
                sld_id_list = prs.slides._sldIdLst
                entries_by_part = {
                    str(prs.part.related_part(entry.rId).partname): entry
                    for entry in sld_id_list
                }
                title_entry = next(
                    entry for entry in sld_id_list
                    if str(prs.part.related_part(entry.rId).partname)
                    == str(prs.slides[0].part.partname)
                )
                ordered_entries = [title_entry]
                for item_index in range(len(measurement_pairs)):
                    for slides in photo_slide_batches:
                        ordered_entries.append(
                            entries_by_part[str(slides[item_index].part.partname)]
                        )
                for entry in list(sld_id_list):
                    sld_id_list.remove(entry)
                for entry in ordered_entries:
                    sld_id_list.append(entry)

                def safe(value):
                    value = re.sub(r"[^\w\s-]", "", str(value), flags=re.UNICODE).strip()
                    return re.sub(r"\s+", "_", value)[:80] or "nomalum"

                netto_value = first.get("netto_f")
                if netto_value is None:
                    netto_value = first.get("netto")
                try:
                    netto_label = format(float(str(netto_value).replace(",", ".")), ".15g")
                except (TypeError, ValueError):
                    netto_label = str(netto_value or "nomalum")
                netto_label = safe(netto_label.replace(".", "_"))
                date_text = report_date.strftime("%Y-%m-%d")
                filename = (
                    f"{report_number}. hisobot_{date_text}_{safe(district)}_"
                    f"{safe(farmer)}_netto_{netto_label}kg.pptx"
                )
                path = os.path.join(output_dir, filename)
                prs.save(path)
                created.append(path)

            print(f"{GREEN}PPTX hisobotlar saqlandi: {len(created)} ta{RESET}")
            for path in created:
                print(f"  ⭐ {path}")
            return created

        for report_number, (group_key, items) in enumerate(
            sorted(
                grouped.items(),
                key=lambda group: (
                    group[0][0], group[0][1],
                    group[0][2] if len(group[0]) > 2 else 0,
                ),
            ),
            start=1,
        ):
            district, farmer = group_key[:2]
            items.sort(key=lambda x: (
                x["first_load"]["in_dt"],
                x["second_load"].get("in_dt") or x["first_load"]["in_dt"],
            ))
            prs = Presentation(template_path) if template_path and os.path.exists(template_path) else Presentation()
            prs.slide_width = Inches(13.333)
            prs.slide_height = Inches(1.1)
            while len(prs.slides):
                rId = prs.slides._sldIdLst[-1].rId
                prs.part.drop_rel(rId)
                del prs.slides._sldIdLst[-1]

            for item in items:
                slide = prs.slides.add_slide(prs.slide_layouts[6])
                bg = slide.background.fill
                bg.solid()
                bg.fore_color.rgb = RGBColor(15, 23, 42)

                def add_text(text, left, top, width, height, size=10, color=(248, 250, 252),
                             bold=False, align=PP_ALIGN.LEFT):
                    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
                    tf = box.text_frame
                    tf.clear()
                    tf.word_wrap = True
                    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
                    p = tf.paragraphs[0]
                    p.alignment = align
                    run = p.add_run()
                    run.text = str(text)
                    run.font.name = "Arial"
                    run.font.size = Pt(size)
                    run.font.bold = bold
                    run.font.color.rgb = RGBColor(*color)
                    return box

                def add_card(left, top, width, height, color):
                    shape = slide.shapes.add_shape(
                        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top),
                        Inches(width), Inches(height)
                    )
                    shape.fill.solid()
                    shape.fill.fore_color.rgb = RGBColor(*color)
                    shape.line.color.rgb = RGBColor(*color)
                    return shape

                first = item["first_load"]
                second = item["second_load"]
                add_card(0.15, 0.25, 1.05, 0.55, (20, 78, 72))
                add_text(district, 0.22, 0.29, 0.9, 0.45, 8, (110, 231, 183), True)
                add_text(first["maskan"], 1.45, 0.25, 3.0, 0.55, 11, bold=True)
                add_card(4.65, 0.25, 0.8, 0.55, (75, 35, 48))
                add_text(first["truck"], 4.72, 0.29, 0.66, 0.45, 10, (252, 165, 165), True, PP_ALIGN.CENTER)

                def details(load, left, color=(248, 250, 252)):
                    lines = [
                        f"Tuman: {load['district']}",
                        f"Kirish: {load['in_str']}",
                        f"Chiqish: {load['out_str']}",
                        f"Netto: {load['netto']} kg",
                        f"Fermer: {load['farmer_name']}",
                    ]
                    add_text("\n".join(lines), left, 0.08, 2.25, 0.88, 8, color)

                details(first, 5.65)
                details(second, 8.25, (248, 250, 252))
                add_card(11.05, 0.25, 1.95, 0.55, (75, 35, 48))
                add_text("1-yuk chiqmasdan 2-yuk kirdi", 11.16, 0.29, 1.73, 0.38, 8, (252, 165, 165), True)
                add_text(f"{item['overlap_minutes']} daqiqa ustma-ust tushgan",
                         11.1, 0.83, 1.85, 0.18, 7, (245, 158, 11), True)

            def safe(value):
                value = re.sub(r"[^\w\s-]", "", str(value), flags=re.UNICODE).strip()
                return re.sub(r"\s+", "_", value)[:80] or "nomalum"

            date_text = report_date.strftime("%Y-%m-%d")
            filename = f"{report_number}. hisobot_{date_text}_{safe(district)}_{safe(farmer)}.pptx"
            path = os.path.join(output_dir, filename)
            prs.save(path)
            created.append(path)

        print(f"{GREEN}PPTX hisobotlar saqlandi: {len(created)} ta{RESET}")
        for path in created:
            print(f"  ⭐ {path}")
        return created

    def export_pptx_rule1_reports(
        self,
        output_dir: str = ".",
        template_path: str = "",
        report_date=None,
        image_headers=None,
    ):
        """1-qoida bo'yicha oraliq kirishlar uchun alohida PPTX hisobotlar yaratadi."""
        report_date = report_date or datetime.now().date()
        synthetic_results = []
        for result in self.r1:
            main_load = result["main_load"]
            if not main_load.get("in_dt") or not self._result_touches_date(result, report_date):
                continue
            for interleaved in result["interleaved_loads"]:
                main_farmer = (main_load.get("farmer_tin") or main_load.get("farmer_name") or "").strip().upper()
                interleaved_farmer = (
                    interleaved.get("farmer_tin")
                    or interleaved.get("farmer_name")
                    or ""
                ).strip().upper()
                if not main_farmer or not interleaved_farmer or main_farmer != interleaved_farmer:
                    continue
                overlap = 0.0
                if main_load.get("out_dt"):
                    overlap = (
                        main_load["out_dt"] - interleaved["in_dt"]
                    ).total_seconds() / 60.0
                synthetic_results.append({
                    "maskan": result["maskan"],
                    "district": result["district"],
                    "truck": main_load["truck"],
                    "first_load": main_load,
                    "second_load": interleaved,
                    "overlap_minutes": round(overlap, 1),
                })

        if not synthetic_results:
            print(f"{GREEN}1-qoida bo'yicha PPTX hisobot topilmadi.{RESET}")
            return []

        result = dict(self.res)
        result["rule3_results"] = synthetic_results
        exporter = ReportGenerator(result)
        created = exporter.export_pptx_reports(
            output_dir=output_dir,
            template_path=template_path,
            report_date=report_date,
            image_headers=image_headers,
        )
        renamed = []
        for path in created:
            source = os.path.basename(path)
            target = os.path.join(
                output_dir,
                source.replace(". hisobot_", ". hisobot_1qoida_", 1),
            )
            os.replace(path, target)
            renamed.append(target)
        print(f"{GREEN}1-qoida PPTX hisobotlar saqlandi: {len(renamed)} ta{RESET}")
        for path in renamed:
            print(f"  ⭐ {path}")
        return renamed

    def export_pptx_rule2_reports(
        self,
        output_dir: str = ".",
        template_path: str = "",
        report_date=None,
        image_headers=None,
    ):
        """2-qoida bo'yicha barcha holatlar, jumladan ayni texnika uchun PPTX yaratadi."""
        report_date = report_date or datetime.now().date()
        synthetic_results = []
        for result in self.r2:
            first_load = result["first_load"]
            if (
                not first_load.get("in_dt")
                or not self._result_touches_date(result, report_date)
            ):
                continue
            synthetic_results.append({
                "maskan": result["maskan"],
                "district": result["district"],
                "truck": first_load.get("truck", ""),
                "first_load": first_load,
                "second_load": result["second_load"],
                "overlap_minutes": result["diff_minutes"],
            })

        if not synthetic_results:
            print(f"{GREEN}2-qoida bo'yicha PPTX hisoboti topilmadi.{RESET}")
            return []

        result = dict(self.res)
        result["rule3_results"] = synthetic_results
        exporter = ReportGenerator(result)
        created = exporter.export_pptx_reports(
            output_dir=output_dir,
            template_path=template_path,
            report_date=report_date,
            image_headers=image_headers,
        )
        renamed = []
        for path in created:
            source = os.path.basename(path)
            target = os.path.join(
                output_dir,
                source.replace(". hisobot_", ". hisobot_2qoida_", 1),
            )
            os.replace(path, target)
            renamed.append(target)
        print(f"{GREEN}2-qoida PPTX hisobotlar saqlandi: {len(renamed)} ta{RESET}")
        for path in renamed:
            print(f"  ⭐ {path}")
        return renamed

    def export_pptx_rule4_reports(
        self,
        output_dir: str = ".",
        template_path: str = "",
        report_date=None,
        image_headers=None,
    ):
        """4-qoida bo'yicha og'irlik va oraliq kirish holatlari uchun PPTX yaratadi."""
        report_date = report_date or datetime.now().date()
        synthetic_results = []
        for item in self.r4:
            first_load = item["first_load"]
            if first_load.get("in_dt") and self._result_touches_date(item, report_date):
                synthetic_results.append({
                    "maskan": item["maskan"],
                    "district": item["district"],
                    "truck": first_load.get("truck", ""),
                    "first_load": first_load,
                    "second_load": item["second_load"],
                    "overlap_minutes": item["overlap_minutes"],
                })
        if not synthetic_results:
            print(f"{GREEN}4-qoida bo'yicha PPTX hisoboti topilmadi.{RESET}")
            return []
        result = dict(self.res)
        result["rule3_results"] = synthetic_results
        created = ReportGenerator(result).export_pptx_reports(
            output_dir=output_dir,
            template_path=template_path,
            report_date=report_date,
            image_headers=image_headers,
        )
        renamed = []
        for path in created:
            target = os.path.join(
                output_dir,
                os.path.basename(path).replace(". hisobot_", ". hisobot_4qoida_", 1),
            )
            os.replace(path, target)
            renamed.append(target)
        print(f"{GREEN}4-qoida PPTX hisobotlar saqlandi: {len(renamed)} ta{RESET}")
        for path in renamed:
            print(f"  ⭐ {path}")
        return renamed

    def export_pptx_rule5_reports(self, output_dir=".", template_path="", report_date=None, image_headers=None):
        report_date = report_date or datetime.now().date()
        synthetic = []
        for item in self.r5:
            first = item["main_load"]
            intermediate_loads = item["interleaved_loads"]
            if (
                not first.get("in_dt")
                or not self._result_touches_date(item, report_date)
                or not intermediate_loads
            ):
                continue

            last_intermediate = max(
                intermediate_loads,
                key=lambda load: load.get("in_dt") or datetime.min,
            )
            synthetic.append({
                "maskan": item["maskan"],
                "district": item["district"],
                "truck": first.get("truck", ""),
                "first_load": first,
                "second_load": last_intermediate,
                "summary_measurements": [first, *intermediate_loads],
                "all_intermediate_loads": intermediate_loads,
                "measurement_numbers": [
                    (first, 0),
                    (last_intermediate, len(intermediate_loads)),
                ],
                "overlap_minutes": round(
                    (first["out_dt"] - last_intermediate["in_dt"]).total_seconds() / 60.0,
                    1,
                ),
            })
        if not synthetic:
            print(f"{GREEN}5-qoida bo'yicha PPTX hisoboti topilmadi.{RESET}")
            return []
        if not template_path or not os.path.isfile(template_path):
            template_display = template_path or "ko'rsatilmagan"
            raise FileNotFoundError(
                f"5-qoida PPTX shabloni topilmadi: {template_display}"
            )
        result = dict(self.res)
        result["rule3_results"] = synthetic
        created = ReportGenerator(result).export_pptx_reports(
            output_dir,
            template_path,
            report_date,
            image_headers,
            one_report_per_case=True,
        )
        renamed = []
        for path in created:
            filename = os.path.basename(path)
            netto_match = re.search(r"(_netto_[^/]+)(?=\.pptx$)", filename)
            netto_suffix = netto_match.group(1) if netto_match else ""
            rule5_filename = filename.replace(
                ". hisobot_", ". hisobot_5qoida_", 1
            )
            if netto_suffix:
                rule5_filename = rule5_filename.replace(netto_suffix, "", 1)
                marker = ". hisobot_5qoida_"
                rule5_filename = rule5_filename.replace(
                    marker, marker + netto_suffix.lstrip("_") + "_", 1
                )
            target = os.path.join(
                output_dir, rule5_filename
            )
            os.replace(path, target)
            renamed.append(target)
        return renamed

    def export_html(self, output_file: str = "monitoring_dashboard.html", report_date=None):
        if report_date is not None:
            def belongs_to_date(result):
                return self._result_touches_date(result, report_date)

            daily_result = dict(self.res)
            daily_result["rule1_results"] = [
                result for result in self.r1 if belongs_to_date(result)
            ]
            daily_result["rule2_results"] = [
                result for result in self.r2 if belongs_to_date(result)
            ]
            daily_result["rule3_results"] = [
                result for result in self.r3 if belongs_to_date(result)
            ]
            daily_result["rule4_results"] = [
                result for result in self.r4 if belongs_to_date(result)
            ]
            day_key = report_date.isoformat()
            daily_result["total_loads"] = self.res.get(
                "daily_load_counts", {}
            ).get(day_key, 0)
            daily_result["total_maskans"] = self.res.get(
                "daily_maskan_counts", {}
            ).get(day_key, 0)
            daily_result["rule1_count"] = len(daily_result["rule1_results"])
            daily_result["rule2_count"] = len(daily_result["rule2_results"])
            daily_result["rule3_count"] = len(daily_result["rule3_results"])
            daily_result["rule4_count"] = len(daily_result["rule4_results"])
            daily_result["rule5_results"] = [
                result for result in self.r5 if belongs_to_date(result)
            ]
            daily_result["rule5_count"] = len(daily_result["rule5_results"])
            daily_reporter = ReportGenerator(daily_result)
            return daily_reporter.export_html(output_file)

        now_str = datetime.now().strftime("%d.%m.%Y %H:%M:%S")
        
        html = f"""<!DOCTYPE html>
<html lang="uz">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Paxta Monitoringi: Anomaliyalar hisoboti</title>
    <style>
        :root {{
            --bg: #0f172a;
            --card-bg: #1e293b;
            --text: #f8fafc;
            --text-muted: #94a3b8;
            --danger: #ef4444;
            --warning: #f59e0b;
            --primary: #3b82f6;
            --success: #10b981;
            --border: #334155;
        }}
        body {{
            font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif;
            background-color: var(--bg);
            color: var(--text);
            margin: 0;
            padding: 24px;
        }}
        .header {{
            display: flex;
            justify-content: space-between;
            align-items: center;
            border-bottom: 1px solid var(--border);
            padding-bottom: 20px;
            margin-bottom: 24px;
        }}
        .header h1 {{ margin: 0; font-size: 24px; color: #fff; }}
        .header p {{ margin: 4px 0 0 0; color: var(--text-muted); font-size: 14px; }}
        .stats-grid {{
            display: grid;
            grid-template-columns: repeat(auto-fit, minmax(200px, 1fr));
            gap: 16px;
            margin-bottom: 28px;
        }}
        .stat-card {{
            background: var(--card-bg);
            border: 1px solid var(--border);
            border-radius: 12px;
            padding: 20px;
        }}
        .stat-card.alert {{
            border-color: var(--danger);
            background: rgba(239, 68, 68, 0.1);
        }}
        .stat-val {{ font-size: 32px; font-weight: bold; margin-top: 6px; }}
        .stat-val.red {{ color: var(--danger); }}
        .stat-val.yellow {{ color: var(--warning); }}
        .stat-val.blue {{ color: var(--primary); }}
        .stat-label {{ font-size: 12px; color: var(--text-muted); text-transform: uppercase; letter-spacing: 0.5px; }}
        
        .tabs {{ display: flex; gap: 10px; margin-bottom: 16px; }}
        .tab-btn {{
            background: var(--card-bg);
            color: var(--text);
            border: 1px solid var(--border);
            padding: 10px 20px;
            border-radius: 8px;
            cursor: pointer;
            font-weight: 600;
        }}
        .tab-btn.active {{ background: var(--primary); border-color: var(--primary); }}
        .tab-btn.alert {{ border-color: var(--danger); color: #fca5a5; }}
        .tab-btn.alert.active {{ background: var(--danger); color: #fff; }}
        .tab-content {{ display: none; }}
        .tab-content.active {{ display: block; }}
        
        table {{
            width: 100%;
            border-collapse: collapse;
            background: var(--card-bg);
            border-radius: 12px;
            overflow: hidden;
            font-size: 13px;
        }}
        th, td {{
            padding: 12px 14px;
            text-align: left;
            border-bottom: 1px solid var(--border);
        }}
        th {{ background: #182234; color: var(--text-muted); font-weight: 600; text-transform: uppercase; font-size: 11px; }}
        tr:hover {{ background: rgba(255,255,255,0.03); }}
        .badge {{
            display: inline-block;
            padding: 4px 8px;
            border-radius: 6px;
            font-size: 11px;
            font-weight: 600;
        }}
        .badge-district {{ background: rgba(16, 185, 129, 0.15); color: #6ee7b7; border: 1px solid rgba(16, 185, 129, 0.3); }}
        .badge-red {{ background: rgba(239, 68, 68, 0.2); color: #fca5a5; border: 1px solid rgba(239, 68, 68, 0.3); }}
        .badge-yellow {{ background: rgba(245, 158, 11, 0.2); color: #fcd34d; border: 1px solid rgba(245, 158, 11, 0.3); }}
        .badge-blue {{ background: rgba(59, 130, 246, 0.2); color: #93c5fd; border: 1px solid rgba(59, 130, 246, 0.3); }}
        .case-toolbar, .sort-toolbar {{
            display: flex;
            align-items: center;
            flex-wrap: wrap;
            gap: 10px;
            background: var(--card-bg);
            border: 1px solid var(--border);
            border-radius: 10px;
            padding: 10px 14px;
            margin: 0 0 16px;
        }}
        .case-toolbar button, .source-link, .sort-toolbar select {{
            display: inline-flex;
            align-items: center;
            justify-content: center;
            min-height: 44px;
            padding: 8px 12px;
            border: 1px solid var(--border);
            border-radius: 8px;
            background: #182234;
            color: #bfdbfe;
            font: inherit;
            text-decoration: none;
            cursor: pointer;
            touch-action: manipulation;
        }}
        .case-toolbar button:hover, .source-link:hover, .sort-toolbar select:hover {{ background: #24344f; }}
        .sort-toolbar select {{ color-scheme: dark; }}
        .case-checkbox {{
            width: 20px;
            height: 20px;
            accent-color: var(--primary);
            cursor: pointer;
            touch-action: manipulation;
        }}
        .source-links {{ display: flex; flex-direction: column; align-items: flex-start; gap: 6px; }}
        .source-link-unavailable {{ color: var(--text-muted); font-size: 12px; }}
        tr.case-selected {{ outline: 2px solid var(--primary); outline-offset: -2px; }}
        @media (max-width: 700px) {{
            body {{ padding: 12px; }}
            .tabs {{ overflow-x: auto; }}
            .tab-btn {{ white-space: nowrap; }}
            table {{ min-width: 900px; }}
            .tab-content {{ overflow-x: auto; }}
        }}
    </style>
</head>
<body>
    <div class="header">
        <div>
            <h1>🌿 Paxta Qabul Qilish Monitoringi: Tumanlar va Maskanlar Kesimida</h1>
            <p>Hosil-qabuli.uz real vaqtdagi tahlil hisoboti | Vaqt: {now_str}</p>
        </div>
    </div>

    <div class="stats-grid">
        <div class="stat-card alert">
            <div class="stat-label">🚨 3-QOIDA: AYLANTIRISH (AYNI TEXNIKA)</div>
            <div class="stat-val red">{self.res['rule3_count']}</div>
        </div>
        <div class="stat-card">
            <div class="stat-label">1-Qoida: Oraliqdagi kirishlar</div>
            <div class="stat-val red">{self.res['rule1_count']}</div>
        </div>
        <div class="stat-card">
            <div class="stat-label">2-Qoida: 15 min ichidagi kirishlar</div>
            <div class="stat-val yellow">{self.res['rule2_count']}</div>
        </div>
        <div class="stat-card">
            <div class="stat-label">4-Qoida: Og'irlik farqi va oraliq kirish</div>
            <div class="stat-val red">{self.res.get('rule4_count', 0)}</div>
        </div>
        <div class="stat-card">
            <div class="stat-label">Jami tekshirilgan yuklar</div>
            <div class="stat-val blue">{self.res['total_loads']}</div>
        </div>
    </div>

    <div class="tabs">
        <button class="tab-btn alert active" onclick="switchTab('tab3', this)">🚨 3-Qoida: Ayni texnika bilan aylantirish ({self.res['rule3_count']})</button>
        <button class="tab-btn" onclick="switchTab('tab1', this)">1-Qoida: Oraliqdagi kirishlar ({self.res['rule1_count']})</button>
        <button class="tab-btn" onclick="switchTab('tab2', this)">2-Qoida: 15 daqiqa ichidagi kirishlar ({self.res['rule2_count']})</button>
        <button class="tab-btn" onclick="switchTab('tab4', this)">4-Qoida: Og'irlik farqi ({self.res.get('rule4_count', 0)})</button>
        <button class="tab-btn" onclick="switchTab('tab5', this)">5-Qoida: Mos fermer va oraliq kirishlar ({self.res.get('rule5_count', 0)})</button>
    </div>
    <div class="case-toolbar" aria-label="Holatlarni tanlash">
        <button type="button" onclick="setAllCases(true)">Barchasini tanlash</button>
        <button type="button" onclick="setAllCases(false)">Tanlovni bekor qilish</button>
        <strong id="selected-case-count" aria-live="polite">0 ta holat tanlangan</strong>
    </div>
    <div class="sort-toolbar" aria-label="Jadvalni saralash">
        <label for="sort-key">Saralash:</label>
        <select id="sort-key" onchange="sortActiveTable()">
            <option value="count">Oraliq kirishlar soni</option>
            <option value="maskan">Maskan</option>
            <option value="farmer">Fermer xo'jaligi</option>
            <option value="netto">Netto</option>
        </select>
        <select id="sort-direction" aria-label="Saralash tartibi" onchange="sortActiveTable()">
            <option value="desc">Kamayish tartibida</option>
            <option value="asc">O'sish tartibida</option>
        </select>
    </div>

    <!-- TAB 3 -->
    <div id="tab3" class="tab-content active">
        <div style="background:rgba(239,68,68,0.15); border:1px solid var(--danger); padding:12px 16px; border-radius:8px; margin-bottom:14px; font-size:13px;">
            <strong>Diqqat!</strong> Texnika hali maskandan chiqib ketmasdan turib, ayni shu davlat raqami bilan yana qayta kirish o'tkazilgan holatlar. Bu 100% soxta o'lchash ("aylantirish") hisoblanadi.
        </div>
        <table>
            <thead>
                <tr>
                    <th>Tanlash</th>
                    <th>Tuman</th>
                    <th>Maskan</th>
                    <th>Texnika Raqami</th>
                    <th>1-Yuk (Asosiy)</th>
                    <th>2-Yuk (Oraliqda qayta kiritilgan)</th>
                    <th>Holati va Tafsilot</th>
                    <th>Manba</th>
                </tr>
            </thead>
            <tbody>
"""
        for idx, item in enumerate(self.r3, start=1):
            f1 = item["first_load"]
            f2 = item["second_load"]
            sort_attrs = self._sort_attributes(item, 1, f1)
            checkbox, source_links = self._case_controls(
                f"rule3-{idx}", (f1, f2)
            )
            html += f"""<tr {sort_attrs}>
                <td>{checkbox}</td>
                <td><span class='badge badge-district'>{item['district']}</span></td>
                <td><strong>{item['maskan']}</strong></td>
                <td><span class='badge badge-red' style='font-size:13px;'>{item['truck']}</span></td>
                <td>
                    Tuman: <strong>{f1['district']}</strong><br>
                    Kirish: <strong>{f1['in_str']}</strong><br>
                    Chiqish: <strong>{f1['out_str']}</strong><br>
                    Netto: {f1['netto']} kg<br>
                    <small>Fermer: {f1['farmer_name']}</small>
                </td>
                <td>
                    Tuman: <strong>{f2['district']}</strong><br>
                    Kirish: <strong style='color:var(--danger)'>{f2['in_str']}</strong><br>
                    Chiqish: <strong>{f2['out_str']}</strong><br>
                    Netto: {f2['netto']} kg<br>
                    <small>Fermer: {f2['farmer_name']}</small>
                </td>
                <td>
                    <span class='badge badge-red'>1-yuk chiqmasdan 2-yuk kirdi</span><br>
                    <small style='color:var(--warning)'>{item['overlap_minutes']} daqiqa ustma-ust tushgan</small>
                </td><td>{source_links}</td>
            </tr>"""

        html += """</tbody>
        </table>
    </div>

    <!-- TAB 1 -->
    <div id="tab1" class="tab-content">
        <table>
            <thead>
                <tr>
                    <th>Tanlash</th>
                    <th>Tuman</th>
                    <th>Maskan</th>
                    <th>Asosiy Texnika</th>
                    <th>Kirish / Chiqish vaqti</th>
                    <th>Asosiy Netto</th>
                    <th>Fermer</th>
                    <th>Oraliqda kirgan boshqa texnikalar</th>
                    <th>Manba</th>
                </tr>
            </thead>
            <tbody>
"""
        for idx, item in enumerate(self.r1, start=1):
            m = item["main_load"]
            sort_attrs = self._sort_attributes(
                item, len(item["interleaved_loads"]), m
            )
            checkbox, source_links = self._case_controls(
                f"rule1-{idx}", (m, *item["interleaved_loads"])
            )
            subs_html = ""
            for sub in item["interleaved_loads"]:
                subs_html += f"<div style='margin-bottom:4px;'><span class='badge badge-district' style='font-size:10px;'>{sub['district']}</span> <span class='badge badge-red'>{sub['truck']}</span> Kirish: {sub['in_str']} | Netto: {sub['netto']} kg (Fermer: {sub['farmer_name']})</div>"

            html += f"""<tr {sort_attrs}>
                <td>{checkbox}</td>
                <td><span class='badge badge-district'>{item['district']}</span></td>
                <td><strong>{item['maskan']}</strong></td>
                <td><span class='badge badge-blue'>{m['truck']}</span></td>
                <td>{m['in_str']}<br><span style='color:var(--text-muted)'>chiqish:</span> {m['out_str']}</td>
                <td><strong>{m['netto']} kg</strong></td>
                <td>{m['farmer_name']}</td>
                <td>{subs_html}</td><td>{source_links}</td>
            </tr>"""

        html += """</tbody>
        </table>
    </div>

    <!-- TAB 2 -->
    <div id="tab2" class="tab-content">
        <table>
            <thead>
                <tr>
                    <th>Tanlash</th>
                    <th>Tuman</th>
                    <th>Maskan</th>
                    <th>Vaqt oralig'i</th>
                    <th>1-Texnika (Kirish)</th>
                    <th>2-Texnika (Kirish)</th>
                    <th>Holati</th>
                    <th>Manba</th>
                </tr>
            </thead>
            <tbody>
"""
        for idx, item in enumerate(self.r2, start=1):
            f1 = item["first_load"]
            f2 = item["second_load"]
            sort_attrs = self._sort_attributes(item, 1, f1)
            checkbox, source_links = self._case_controls(
                f"rule2-{idx}", (f1, f2)
            )
            badge = "<span class='badge badge-red'>⚠️ AYNI TEXNIKA</span>" if item["is_same_truck"] else "<span class='badge badge-yellow'>Tezkor kirish</span>"
            
            html += f"""<tr {sort_attrs}>
                <td>{checkbox}</td>
                <td><span class='badge badge-district'>{item['district']}</span></td>
                <td><strong>{item['maskan']}</strong></td>
                <td><strong style='color:var(--warning)'>{item['diff_minutes']} daqiqa</strong></td>
                <td><span class='badge badge-blue'>{f1['truck']}</span> {f1['in_str']}<br><small>[{f1['district']}] {f1['farmer_name']} ({f1['netto']} kg)</small></td>
                <td><span class='badge badge-blue'>{f2['truck']}</span> {f2['in_str']}<br><small>[{f2['district']}] {f2['farmer_name']} ({f2['netto']} kg)</small></td>
                <td>{badge}</td><td>{source_links}</td>
            </tr>"""

        html += """</tbody>
        </table>
    </div>

    <!-- TAB 4 -->
    <div id="tab4" class="tab-content">
        <table>
            <thead>
                <tr>
                    <th>Tanlash</th>
                    <th>Maskan tumani</th><th>Maskan</th><th>Og'irlik farqi</th>
                    <th>1-yuk fermer xo'jaligi</th><th>1-yuk fermer tumani</th>
                    <th>1-yuk</th><th>2-yuk fermer xo'jaligi</th>
                    <th>2-yuk fermer tumani</th><th>Oraliqdagi boshqa yuk</th>
                    <th>Manba</th>
                </tr>
            </thead>
            <tbody>
"""
        for idx, item in enumerate(self.r4, start=1):
            first = item["first_load"]
            second = item["second_load"]
            sort_attrs = self._sort_attributes(item, 1, first)
            checkbox, source_links = self._case_controls(
                f"rule4-{idx}", (first, second)
            )
            html += f"""<tr {sort_attrs}>
                <td>{checkbox}</td>
                <td><span class='badge badge-district'>{item['district']}</span></td>
                <td><strong>{item['maskan']}</strong></td>
                <td><strong style='color:var(--danger)'>{item['weight_difference']} kg</strong></td>
                <td><strong>{first['farmer_name']}</strong></td>
                <td><span class='badge badge-district'>{first['district']}</span></td>
                <td><span class='badge badge-blue'>{first['truck']}</span>
                    {first['in_str']} → {first['out_str']}<br>
                    Brutto: {first['fullweight']} kg | Tara: {first['taraweight']} kg | Netto: {first['netto']} kg<br>
                    Farqlar: brutto {item['weight_differences']['brutto']} kg,
                    tara {item['weight_differences']['tara']} kg,
                    netto {item['weight_differences']['netto']} kg</td>
                <td><strong>{second['farmer_name']}</strong></td>
                <td><span class='badge badge-district'>{second['district']}</span></td>
                <td><span class='badge badge-red'>{second['truck']}</span>
                    {second['in_str']}</td><td>{source_links}</td>
            </tr>"""
        html += """</tbody>
        </table>
    </div>

    <div id="tab5" class="tab-content">
        <table><thead><tr>
            <th>Tanlash</th>
            <th>Tuman</th><th>Maskan</th><th>Asosiy texnika</th><th>Kirish / chiqish vaqti</th>
            <th>Asosiy netto</th><th>Fermer xo'jaligi</th><th>Og'irlik shartiga mos oraliqda kirgan boshqa texnikalar</th>
            <th>Manba</th>
        </tr></thead><tbody>
"""
        for idx, item in enumerate(self.r5, start=1):
            main = item["main_load"]
            sort_attrs = self._sort_attributes(
                item, len(item["interleaved_loads"]), main
            )
            checkbox, source_links = self._case_controls(
                f"rule5-{idx}", (main, *item["interleaved_loads"])
            )
            subs_html = ""
            for sub in item["interleaved_loads"]:
                differences = self._weight_differences(main, sub) or {}
                subs_html += (
                    f"<div style='margin-bottom:6px;'>"
                    f"<span class='badge badge-district'>{sub['district']}</span> "
                    f"<span class='badge badge-red'>{sub['truck']}</span> "
                    f"Kirish: {sub['in_str']} | Chiqish: {sub['out_str']} | "
                    f"Netto: {sub['netto']} kg<br>"
                    f"<small>Fermer: {sub['farmer_name']} | "
                    f"Brutto farqi: {differences.get('brutto', 0):.1f} kg, "
                    f"Tara farqi: {differences.get('tara', 0):.1f} kg, "
                    f"Netto farqi: {differences.get('netto', 0):.1f} kg</small></div>"
                )
            html += f"""<tr {sort_attrs}>
                <td>{checkbox}</td>
                <td><span class='badge badge-district'>{item['district']}</span></td>
                <td><strong>{item['maskan']}</strong></td>
                <td><span class='badge badge-blue'>{main['truck']}</span></td>
                <td>{main['in_str']}<br><span style='color:var(--text-muted)'>chiqish:</span> {main['out_str']}</td>
                <td><strong>{main['netto']} kg</strong></td>
                <td>{main['farmer_name']}</td>
                <td>{subs_html}</td><td>{source_links}</td>
            </tr>"""
        html += """</tbody></table>
    </div>

    <script>
        function switchTab(tabId, selectedButton) {
            document.querySelectorAll('.tab-btn').forEach(b => b.classList.remove('active'));
            document.querySelectorAll('.tab-content').forEach(c => c.classList.remove('active'));
            selectedButton.classList.add('active');
            document.getElementById(tabId).classList.add('active');
            sortActiveTable();
        }
        function sortActiveTable() {
            const tableBody = document.querySelector('.tab-content.active table tbody');
            if (!tableBody) return;
            const key = document.getElementById('sort-key').value;
            const direction = document.getElementById('sort-direction').value === 'asc' ? 1 : -1;
            const attribute = `data-sort-${key}`;
            const collator = new Intl.Collator('uz', {
                sensitivity: 'base',
                numeric: true
            });
            const rows = [...tableBody.rows];
            rows.sort((a, b) => {
                const left = a.getAttribute(attribute) || '';
                const right = b.getAttribute(attribute) || '';
                if (!left && right) return 1;
                if (left && !right) return -1;
                let comparison;
                if (key === 'count' || key === 'netto') {
                    comparison = (Number(left) || 0) - (Number(right) || 0);
                } else {
                    comparison = collator.compare(left, right);
                }
                return comparison * direction;
            });
            rows.forEach(row => tableBody.appendChild(row));
        }
        function updateSelectedCases() {
            const checkboxes = [...document.querySelectorAll('.case-checkbox')];
            const selected = checkboxes.filter(input => input.checked);
            document.getElementById('selected-case-count').textContent =
                `${selected.length} ta holat tanlangan`;
            checkboxes.forEach(input => {
                const row = input.closest('tr');
                if (row) row.classList.toggle('case-selected', input.checked);
            });
        }
        function setAllCases(checked) {
            document.querySelectorAll('.case-checkbox').forEach(input => {
                input.checked = checked;
            });
            updateSelectedCases();
        }
        document.addEventListener('change', event => {
            if (event.target.matches('.case-checkbox')) updateSelectedCases();
        });
    </script>
</body>
</html>"""

        with open(output_file, "w", encoding="utf-8") as fp:
            fp.write(html)
        print(f"{GREEN}Interaktiv HTML dashboard saqlandi: {output_file}{RESET}")
        return output_file
